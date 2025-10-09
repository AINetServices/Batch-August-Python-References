"""
Custom tools for resume processing and text extraction.
"""

import os
import io
<<<<<<< HEAD
from typing import Any, Dict, List, Type
=======
from typing import Any, Dict, List
>>>>>>> 5ef3108f3d16761ce7924e7b6ff831895e47f517
from langchain.tools import BaseTool
from langchain_groq import ChatGroq
import PyPDF2
import docx
<<<<<<< HEAD
from pydantic import BaseModel, Field


# Define input schemas for each tool
class ResumeParserInput(BaseModel):
    file_path: str = Field(description="Path to the resume file to parse")


class ReferenceExtractorInput(BaseModel):
    resume_text: str = Field(description="The resume text to extract references from")
    role: str = Field(default="", description="Target role for reference extraction")


class ResumeAnalyzerInput(BaseModel):
    resume_text: str = Field(description="The resume text to analyze")
    target_role: str = Field(default="", description="Target role for analysis")


# Create simple function-based tools
def create_resume_parser_tool(llm: ChatGroq):
    """Create a resume parser tool"""
    
    def parse_resume(file_path: str) -> str:
        """Parse a resume file and return extracted text"""
        try:
            with open(file_path, 'rb') as file:
                filename = file_path.lower()
                if filename.endswith('.pdf'):
                    return _extract_pdf_text(file)
                elif filename.endswith(('.docx', '.doc')):
                    return _extract_docx_text(file)
                else:
                    return "Unsupported file format"
        except Exception as e:
            return f"Error parsing file: {str(e)}"
    
    return parse_resume


def create_reference_extractor_tool(llm: ChatGroq):
    """Create a reference extractor tool"""
    
    def extract_references(resume_text: str, role: str = "") -> Dict[str, Any]:
=======
from pydantic import BaseModel

class ResumeParserTool(BaseTool):
    """Tool for parsing resume files (PDF, DOCX) and extracting text"""
    
    name = "resume_parser"
    description = "Parse resume files and extract text content"
    
    def __init__(self, llm: ChatGroq):
        super().__init__()
        self.llm = llm

    def _run(self, file_path: str) -> str:
        """Parse a resume file and return extracted text"""
        try:
            with open(file_path, 'rb') as file:
                return self.parse_file_content(file, file_path)
        except Exception as e:
            return f"Error parsing file: {str(e)}"

    def parse_file_content(self, file_content: io.BytesIO, filename: str) -> str:
        """Parse file content based on file extension"""
        try:
            if filename.lower().endswith('.pdf'):
                return self._extract_pdf_text(file_content)
            elif filename.lower().endswith(('.docx', '.doc')):
                return self._extract_docx_text(file_content)
            else:
                return "Unsupported file format"
        except Exception as e:
            return f"Error extracting text: {str(e)}"

    def _extract_pdf_text(self, file_content: io.BytesIO) -> str:
        """Extract text from PDF file"""
        try:
            pdf_reader = PyPDF2.PdfReader(file_content)
            text = ""
            
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            
            return text.strip()
        except Exception as e:
            return f"Error reading PDF: {str(e)}"

    def _extract_docx_text(self, file_content: io.BytesIO) -> str:
        """Extract text from DOCX file"""
        try:
            doc = docx.Document(file_content)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
        except Exception as e:
            return f"Error reading DOCX: {str(e)}"

    async def _arun(self, file_path: str) -> str:
        """Async version of _run"""
        return self._run(file_path)


class ReferenceExtractorTool(BaseTool):
    """Tool for extracting reference information from resume text"""
    
    name = "reference_extractor"
    description = "Extract reference contacts and professional relationships from resume text"
    
    def __init__(self, llm: ChatGroq):
        super().__init__()
        self.llm = llm

    def _run(self, resume_text: str, role: str = "") -> Dict[str, Any]:
>>>>>>> 5ef3108f3d16761ce7924e7b6ff831895e47f517
        """Extract references from resume text"""
        try:
            prompt = f"""
            Analyze this resume text and extract professional reference information:

            Resume Text:
            {resume_text}

            Target Role: {role}

<<<<<<< HEAD
            Extract and identify references and return as JSON.
            """
            
            response = llm.invoke(prompt)
            return _parse_llm_response(response.content)
        except Exception as e:
            return {"error": f"Failed to extract references: {str(e)}"}
    
    return extract_references


def create_resume_analyzer_tool(llm: ChatGroq):
    """Create a resume analyzer tool"""
    
    def analyze_resume(resume_text: str, target_role: str = "") -> Dict[str, Any]:
=======
            Extract and identify:
            1. Names of potential references (supervisors, managers, colleagues)
            2. Contact information (email addresses, phone numbers if available)
            3. Company/organization names
            4. Professional relationships and context
            5. Duration of working relationships
            6. Specific projects or achievements mentioned together

            Focus on:
            - Direct supervisors and managers
            - Close colleagues and team members
            - Clients or stakeholders (if mentioned positively)
            - Anyone mentioned in a professional context

            Return structured information as JSON with the following format:
            {{
                "references": [
                    {{
                        "name": "Full Name",
                        "email": "email@example.com",
                        "company": "Company Name",
                        "relationship": "Direct Supervisor/Colleague/etc",
                        "years_worked": "2018-2020",
                        "context": "Brief description of working relationship",
                        "confidence": "high/medium/low"
                    }}
                ],
                "total_found": number,
                "extraction_notes": "Any important notes about the extraction process"
            }}
            """
            
            response = self.llm.invoke(prompt)
            
            # Parse the response and extract structured data
            return self._parse_llm_response(response.content)
            
        except Exception as e:
            return {"error": f"Failed to extract references: {str(e)}"}

    def _parse_llm_response(self, response: str) -> Dict[str, Any]:
        """Parse LLM response into structured format"""
        try:
            import json
            
            # Find JSON in the response
            start_idx = response.find('{')
            end_idx = response.rfind('}') + 1
            
            if start_idx >= 0 and end_idx > start_idx:
                json_str = response[start_idx:end_idx]
                return json.loads(json_str)
            else:
                return {
                    "references": [],
                    "total_found": 0,
                    "extraction_notes": "Failed to parse LLM response",
                    "raw_response": response
                }
        except json.JSONDecodeError:
            return {
                "references": [],
                "total_found": 0,
                "extraction_notes": "Invalid JSON in LLM response",
                "raw_response": response
            }

    async def _arun(self, resume_text: str, role: str = "") -> Dict[str, Any]:
        """Async version of _run"""
        return self._run(resume_text, role)


class ResumeAnalyzerTool(BaseTool):
    """Advanced tool for analyzing resume content and extracting insights"""
    
    name = "resume_analyzer"
    description = "Analyze resume content for skills, experience, and professional background"
    
    def __init__(self, llm: ChatGroq):
        super().__init__()
        self.llm = llm

    def _run(self, resume_text: str, target_role: str = "") -> Dict[str, Any]:
>>>>>>> 5ef3108f3d16761ce7924e7b6ff831895e47f517
        """Analyze resume content and extract comprehensive insights"""
        try:
            analysis_prompt = f"""
            Perform a comprehensive analysis of this resume for the target role: {target_role}

            Resume Text:
            {resume_text}

<<<<<<< HEAD
            Analyze and return as structured JSON.
            """
            
            response = llm.invoke(analysis_prompt)
            return _parse_analysis_response(response.content)
        except Exception as e:
            return {"error": f"Failed to analyze resume: {str(e)}"}
    
    return analyze_resume


# Helper functions
def _extract_pdf_text(file_content: io.BytesIO) -> str:
    """Extract text from PDF file"""
    try:
        pdf_reader = PyPDF2.PdfReader(file_content)
        text = ""
        
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        
        return text.strip()
    except Exception as e:
        return f"Error reading PDF: {str(e)}"


def _extract_docx_text(file_content: io.BytesIO) -> str:
    """Extract text from DOCX file"""
    try:
        doc = docx.Document(file_content)
        text = ""
        
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
        
        return text.strip()
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"


def _parse_llm_response(response: str) -> Dict[str, Any]:
    """Parse LLM response into structured format"""
    try:
        import json
        
        # Find JSON in the response
        start_idx = response.find('{')
        end_idx = response.rfind('}') + 1
        
        if start_idx >= 0 and end_idx > start_idx:
            json_str = response[start_idx:end_idx]
            return json.loads(json_str)
        else:
            return {
                "references": [],
                "total_found": 0,
                "extraction_notes": "Failed to parse LLM response",
                "raw_response": response
            }
    except json.JSONDecodeError:
        return {
            "references": [],
            "total_found": 0,
            "extraction_notes": "Invalid JSON in LLM response",
            "raw_response": response
        }


def _parse_analysis_response(response: str) -> Dict[str, Any]:
    """Parse comprehensive analysis response"""
    try:
        import json
        
        start_idx = response.find('{')
        end_idx = response.rfind('}') + 1
        
        if start_idx >= 0 and end_idx > start_idx:
            json_str = response[start_idx:end_idx]
            return json.loads(json_str)
        else:
            return {
                "error": "Failed to parse analysis response",
                "raw_response": response
            }
    except json.JSONDecodeError:
        return {
            "error": "Invalid JSON in analysis response",
            "raw_response": response
        }


# Simple wrapper classes
class ResumeParserTool:
    def __init__(self, llm: ChatGroq):
        self.llm = llm
        self._parse_func = create_resume_parser_tool(llm)
    
    def _run(self, file_path: str) -> str:
        return self._parse_func(file_path)


class ReferenceExtractorTool:
    def __init__(self, llm: ChatGroq):
        self.llm = llm
        self._extract_func = create_reference_extractor_tool(llm)
    
    def _run(self, resume_text: str, role: str = "") -> Dict[str, Any]:
        return self._extract_func(resume_text, role)


class ResumeAnalyzerTool:
    def __init__(self, llm: ChatGroq):
        self.llm = llm
        self._analyze_func = create_resume_analyzer_tool(llm)
    
    def _run(self, resume_text: str, target_role: str = "") -> Dict[str, Any]:
        return self._analyze_func(resume_text, target_role)
=======
            Analyze and extract:

            1. APPLICANT PROFILE:
               - Full name
               - Contact information
               - Current position/title
               - Years of experience
               - Location

            2. PROFESSIONAL EXPERIENCE:
               - Companies worked at
               - Positions held
               - Employment dates
               - Key achievements and responsibilities

            3. SKILLS AND COMPETENCIES:
               - Technical skills
               - Soft skills
               - Certifications
               - Education background

            4. REFERENCE INDICATORS:
               - Names mentioned in professional context
               - Supervisors/managers referenced
               - Colleagues or team members
               - Clients or stakeholders
               - Professional relationships described

            5. ROLE ALIGNMENT:
               - How well the background matches {target_role}
               - Relevant experience for the target role
               - Potential strengths and gaps

            Return as structured JSON:
            {{
                "applicant_profile": {{...}},
                "experience": [...],
                "skills": {{...}},
                "potential_references": [...],
                "role_alignment": {{...}},
                "analysis_summary": "Brief summary of findings"
            }}
            """
            
            response = self.llm.invoke(analysis_prompt)
            return self._parse_analysis_response(response.content)
            
        except Exception as e:
            return {"error": f"Failed to analyze resume: {str(e)}"}

    def _parse_analysis_response(self, response: str) -> Dict[str, Any]:
        """Parse comprehensive analysis response"""
        try:
            import json
            
            start_idx = response.find('{')
            end_idx = response.rfind('}') + 1
            
            if start_idx >= 0 and end_idx > start_idx:
                json_str = response[start_idx:end_idx]
                return json.loads(json_str)
            else:
                return {
                    "error": "Failed to parse analysis response",
                    "raw_response": response
                }
        except json.JSONDecodeError:
            return {
                "error": "Invalid JSON in analysis response",
                "raw_response": response
            }

    async def _arun(self, resume_text: str, target_role: str = "") -> Dict[str, Any]:
        """Async version of _run"""
        return self._run(resume_text, target_role)
>>>>>>> 5ef3108f3d16761ce7924e7b6ff831895e47f517
